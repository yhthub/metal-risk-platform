"""Bounded, stateless table conversion; no imputation or medical-unit inference."""
from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
import json
from pathlib import PurePosixPath
import re
from typing import Any, Literal
from zipfile import BadZipFile, ZipFile

import pandas as pd
import numpy as np
from pydantic import BaseModel, ConfigDict, Field, StrictBool, StrictInt, StrictStr, ValidationError, field_validator, model_validator

from .data_service import OUTCOME_CONFIGS, _as_numeric, _get_required_columns


MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_EXPANDED_BYTES = 128 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 300
MAX_TOTAL_CELLS = 2_000_000
MAX_CELL_CHARS = 32_767
MAX_LOGICAL_TEXT_BYTES = 64 * 1024 * 1024
MAX_OUTPUT_BYTES = 80 * 1024 * 1024
PREVIEW_ROWS = 20
MAX_OPTIONS_BYTES = 64 * 1024
# Mirrors the CURRENT _build_analysis cohort rules, not new medical cutoffs.
# The regression suite compares this checker with that analysis implementation.
_SYSTEM_BP_RANGES = {"sbp": (50, 300), "dbp": (30, 200)}
_SYSTEM_MIN_ADULT_AGE = 18
# NHANES 2017–2018 official codebooks (DEMO_J / SMQ_J / ALQ_J):
# https://wwwn.cdc.gov/Nchs/Data/Nhanes/Public/2017/DataFiles/DEMO_J.htm
# https://wwwn.cdc.gov/Nchs/Data/Nhanes/Public/2017/DataFiles/SMQ_J.htm
# https://wwwn.cdc.gov/Nchs/Data/Nhanes/Public/2017/DataFiles/ALQ_J.htm
# These are FIELD-SPECIFIC. RIDRETH3=7 is a legitimate category, not missing.
_NHANES_EXPLICIT_NONRESPONSE = {
    "DEMO_J__DMDEDUC2": {7, 9},
    "DEMO_J__DMDMARTL": {77, 99},
    "SMQ_J__SMQ020": {7, 9},
    "ALQ_J__ALQ111": {7, 9},
}
_NHANES_RESPONSE_CATEGORIES = {
    "DEMO_J__RIAGENDR": {1, 2},
    "DEMO_J__RIDRETH3": {1, 2, 3, 4, 6, 7},
    "DEMO_J__DMDEDUC2": {1, 2, 3, 4, 5},
    "DEMO_J__DMDMARTL": {1, 2, 3, 4, 5, 6},
    "SMQ_J__SMQ020": {1, 2},
    "ALQ_J__ALQ111": {1, 2},
}


class CleaningError(ValueError):
    def __init__(self, message: str, status_code: int = 422):
        super().__init__(message)
        self.status_code = status_code


class ColumnMapping(BaseModel):
    model_config = ConfigDict(extra="forbid")
    source: StrictStr = Field(min_length=1, max_length=256)
    target: StrictStr = Field(min_length=1, max_length=256)

    @field_validator("source", "target")
    @classmethod
    def meaningful_name(cls, value: str) -> str:
        if not value.strip() or "\x00" in value:
            raise ValueError("列名不能为空或包含空字符")
        return value.strip()


class RowFilter(BaseModel):
    model_config = ConfigDict(extra="forbid")
    column: StrictStr = Field(min_length=1, max_length=256)
    operator: Literal["eq", "ne", "contains", "gt", "gte", "lt", "lte", "is_empty", "not_empty"]
    value: StrictStr = Field(default="", max_length=1024)

    @field_validator("column")
    @classmethod
    def meaningful_name(cls, value: str) -> str:
        if not value.strip() or "\x00" in value:
            raise ValueError("筛选列不能为空")
        return value.strip()


class CleaningOptions(BaseModel):
    model_config = ConfigDict(extra="forbid")
    mode: Literal["system", "custom"] = "system"
    outcome_id: StrictStr = Field(default="blood_pressure", min_length=1, max_length=80)
    columns: list[ColumnMapping] = Field(default_factory=list, max_length=MAX_COLUMNS)
    trim_text: StrictBool = True
    drop_blank_rows: StrictBool = True
    drop_duplicates: StrictBool = False
    row_start: StrictInt = Field(default=1, ge=1, le=MAX_ROWS)
    row_end: StrictInt | None = Field(default=None, ge=1, le=MAX_ROWS)
    filters: list[RowFilter] = Field(default_factory=list, max_length=20)
    confirm_system_semantics: StrictBool = False
    drop_invalid_system_rows: StrictBool = False
    require_complete_system_rows: StrictBool = False

    @model_validator(mode="after")
    def unambiguous_columns_and_range(self):
        if self.mode != "system" and (self.drop_invalid_system_rows or self.require_complete_system_rows):
            raise ValueError("不合格记录剔除/字段完整选项仅用于系统格式模式")
        if self.row_end is not None and self.row_end < self.row_start:
            raise ValueError("结束行必须大于或等于起始行")
        sources = [item.source for item in self.columns]
        targets = [item.target for item in self.columns]
        if len(sources) != len(set(sources)):
            raise ValueError("同一来源列不能被重复映射")
        if len(targets) != len(set(targets)):
            raise ValueError("输出列名不能重复")
        return self


def parse_options(options_json: str) -> CleaningOptions:
    if len(options_json.encode("utf-8")) > MAX_OPTIONS_BYTES:
        raise CleaningError("清洗参数超过 64 KB 限制")
    try:
        payload = json.loads(options_json)
    except (ValueError, RecursionError) as exc:
        raise CleaningError("清洗参数必须是合法 JSON 对象") from exc
    try:
        return CleaningOptions.model_validate(payload)
    except ValidationError as exc:
        # Do not include uploaded health data/filter values in validation errors.
        messages = [".".join(map(str, error["loc"])) + ": " + error["msg"]
                    for error in exc.errors(include_input=False)[:5]]
        raise CleaningError("清洗参数不符合要求：" + "；".join(messages)) from exc


def get_cleaning_schema(outcome_id: str = "blood_pressure") -> dict[str, Any]:
    config = OUTCOME_CONFIGS.get(outcome_id)
    if config is None:
        raise CleaningError("不支持的分析结局")
    required = _get_required_columns(config)
    field_info = {
        config["identifier_column"]: {"label": "参与者 ID", "unit": "文本标识符"},
    }
    for kind, columns in config["outcome_columns"].items():
        outcome = next((item for item in config["outcomes"] if kind in item["key"]), {})
        label = {"sbp": "收缩压", "dbp": "舒张压"}.get(kind, kind)
        for index, column in enumerate(columns, start=1):
            field_info[column] = {"label": f"{label}读数 {index}", "unit": outcome.get("unit", "")}
    for field in config["exposures"] + config["confounders"]:
        field_info[field["column"]] = {"label": field["label"], "unit": field.get("unit", "")}
    fields = [{"column": column, **field_info.get(column, {"label": column, "unit": ""}), "required": True}
              for column in required]
    fields.extend({"column": column, "label": "抽样设计（可选）", "unit": "原始编码/权重", "required": False}
                  for column in config.get("survey_design_columns", []) if column not in required)
    return {
        "outcome_id": outcome_id, "name": config["name"], "fields": fields,
        "required_columns": required,
        "notes": [
            "一行代表一位参与者，第一行为列名。六个血压读数列名均需存在；进入现有分析仍须年满 18 岁、收缩压和舒张压各至少 2 次有效读数。",
            "金属指标按现有血液样本及列示单位匹配；不自动换算血液/尿液、浓度单位或分类编码。",
            "分类字段须保持原系统 NHANES 编码含义，不能把男/女、是/否等文本仅改列名后当作数字编码。",
            "缺失单元格保留空白，不补 0、不插补、不推断；是否适合分析由后续群体导入另行检查。",
            "系统就绪表示必需列齐全、已明确确认样本/单位/编码，且保留记录通过当前系统规则检查；不代表临床有效、统计模型已训练或原始数据已验证。",
            "CSV 按 UTF-8 或 GB18030 读取；XLSX 仅读取第一个工作表。Excel 数字单元格不能恢复仅由显示格式产生的前导零，标识符请存为文本。",
            "另有限制为 200 万数据单元格、单元格文本 32,767 字符，宽表的可上传行数因此可能低于 10 万行。",
            "重复引用的 XLSX 共享字符串也计入逻辑文本预算：最多 64 MiB；下载 CSV 最多 80 MiB。",
            "系统行检查沿用当前分析读数范围：收缩压 50–300、舒张压 30–200 mmHg，各至少 2 次有效读数，年龄至少 18 岁；这是现有分析队列规则，不是新的临床正常值。",
            "非空血压值必须为有限正数，暴露/混杂字段的非空值必须为有限数字；不按暴露高值或未知分类编码猜测删除记录。‘字段完整且排除明确非回答’为可选项，默认允许暴露/混杂字段保留原有空白。",
            "字段完整选项仅额外排除教育 7/9、婚姻 77/99、SMQ020 7/9、ALQ111 7/9 等官方明确非回答；RIDRETH3=7 为合法种族类别。未知有限分类编码仅提示、不自动映射或剔除，这不是完整病例或编码有效性认证。",
        ],
        "limits": {"max_file_bytes": MAX_FILE_BYTES, "max_rows": MAX_ROWS, "max_columns": MAX_COLUMNS,
                   "max_total_cells": MAX_TOTAL_CELLS, "max_cell_chars": MAX_CELL_CHARS,
                   "max_logical_text_bytes": MAX_LOGICAL_TEXT_BYTES, "max_output_bytes": MAX_OUTPUT_BYTES,
                   "preview_rows": PREVIEW_ROWS, "accepted_extensions": [".csv", ".xlsx"]},
    }


@dataclass
class Table:
    filename: str
    data: pd.DataFrame
    warnings: list[str]


def _filename(name: str) -> str:
    basename = PurePosixPath(name.replace("\\", "/")).name
    safe = re.sub(r"[^0-9A-Za-z._\-\u4e00-\u9fff]", "_", basename)
    return (safe[:180] or "dataset.csv")


def _headers(raw: list[str], warnings: list[str]) -> list[str]:
    if not raw or len(raw) > MAX_COLUMNS:
        raise CleaningError(f"表格须有列名且不超过 {MAX_COLUMNS} 列")
    headers = [value.strip().lstrip("\ufeff") for value in raw]
    if any(not value or len(value) > 256 or "\x00" in value for value in headers):
        raise CleaningError("列名不能为空、超过 256 字符或包含空字符")
    if len(headers) != len(set(headers)):
        raise CleaningError("原始文件存在重复列名（含去除两端空白后的重复），请先区分列名")
    if headers != raw:
        warnings.append("已去除列名两端空白/BOM；字段映射使用预览中显示的列名。")
    return headers


def _text_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def _csv_table(content: bytes, warnings: list[str]) -> pd.DataFrame:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = content.decode(encoding)
            if encoding == "gb18030":
                warnings.append("原 CSV 按 GB18030 解码；下载统一为带 BOM 的 UTF-8 CSV。")
            break
        except UnicodeDecodeError:
            continue
    else:
        raise CleaningError("CSV 编码无法识别，请另存为 UTF-8 CSV")
    if "\x00" in text:
        raise CleaningError("CSV 含空字符，可能不是受支持的文本编码；请另存为 UTF-8 CSV")
    try:
        reader = csv.reader(StringIO(text, newline=""), strict=True)
        raw_header = next(reader, None)
        if raw_header is None:
            raise CleaningError("文件为空")
        headers = _headers(raw_header, warnings)
        rows = []
        short_rows = 0
        logical_bytes = sum(len(value.encode("utf-8")) for value in headers)
        for row in reader:
            if len(rows) >= MAX_ROWS:
                raise CleaningError(f"文件超过 {MAX_ROWS:,} 行限制")
            if (len(rows) + 1) * len(headers) > MAX_TOTAL_CELLS:
                raise CleaningError(f"表格超过 {MAX_TOTAL_CELLS:,} 个数据单元格预算，请分批上传或减少列数")
            if len(row) > len(headers):
                raise CleaningError(f"第 {len(rows) + 1} 个数据记录的单元格数量超过列名数量；请检查 CSV 分隔符或引号")
            if any(len(value) > MAX_CELL_CHARS for value in row):
                raise CleaningError(f"单元格文本超过 {MAX_CELL_CHARS:,} 字符限制")
            logical_bytes += sum(len(value.encode("utf-8")) for value in row)
            if logical_bytes > MAX_LOGICAL_TEXT_BYTES:
                raise CleaningError("表格逻辑文本总量超过 64 MiB，请减少行列或分批上传")
            if len(row) < len(headers):
                short_rows += 1
                row += [""] * (len(headers) - len(row))
            rows.append(row)
        if short_rows:
            warnings.append(f"{short_rows} 行缺少尾部单元格，保留为空白（未补值）。")
        return pd.DataFrame(rows, columns=headers, dtype=str)
    except csv.Error as exc:
        raise CleaningError("CSV 语法不正确或单元格过长，请使用逗号分隔并正确引用含逗号/换行的文本") from exc


def _xlsx_table(content: bytes, warnings: list[str]) -> pd.DataFrame:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 3000 or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES:
                raise CleaningError("XLSX 展开内容过大，限制为 128 MB / 3000 个内部文件")
            if any(entry.flag_bits & 1 for entry in entries):
                raise CleaningError("不支持加密的工作簿，请上传不带密码的 XLSX")
            if any(entry.filename.lower().endswith("vbaproject.bin") for entry in entries):
                raise CleaningError("不支持带宏的工作簿，请另存为不含宏的 XLSX")
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            if not workbook.worksheets:
                raise CleaningError("工作簿没有可读取工作表")
            sheet = workbook.worksheets[0]
            # Third-party exporters may under/overstate <dimension>. Trusting
            # it silently truncates real rows. Read actual XML records instead;
            # cumulative checks below bound sparse rows and dishonest dimensions.
            sheet.reset_dimensions()
            iterator = sheet.iter_rows(values_only=True)
            raw_header = next(iterator, None)
            if raw_header is None:
                raise CleaningError("第一个工作表为空")
            headers = _headers([_text_cell(value) for value in raw_header], warnings)
            rows = []
            logical_bytes = sum(len(value.encode("utf-8")) for value in headers)
            for row in iterator:
                if len(rows) >= MAX_ROWS or len(row) > MAX_COLUMNS:
                    raise CleaningError("工作表超出行/列限制")
                if (len(rows) + 1) * len(headers) > MAX_TOTAL_CELLS:
                    raise CleaningError(f"工作表超过 {MAX_TOTAL_CELLS:,} 个数据单元格预算")
                values = [_text_cell(value) for value in row]
                if any(len(value) > MAX_CELL_CHARS for value in values):
                    raise CleaningError(f"单元格文本超过 {MAX_CELL_CHARS:,} 字符限制")
                logical_bytes += sum(len(value.encode("utf-8")) for value in values)
                if logical_bytes > MAX_LOGICAL_TEXT_BYTES:
                    raise CleaningError("工作表逻辑文本总量超过 64 MiB（共享字符串重复引用亦计入），请减少行列或分批上传")
                if len(values) > len(headers):
                    raise CleaningError("工作表数据列数超过列名数量")
                values += [""] * (len(headers) - len(values))
                rows.append(values)
            warnings.append("XLSX 仅读取第一个工作表；公式不计算，按原公式文本保留并在下载时防护。")
            return pd.DataFrame(rows, columns=headers, dtype=str)
        finally:
            workbook.close()
    except CleaningError:
        raise
    except (BadZipFile, ImportError, ValueError, KeyError, OSError, TypeError, SyntaxError) as exc:
        raise CleaningError("无法读取 XLSX，请确认文件未损坏、未加密且首行有列名") from exc


def read_table(content: bytes, filename: str) -> Table:
    if not content:
        raise CleaningError("文件为空")
    if len(content) > MAX_FILE_BYTES:
        raise CleaningError("文件超过 25 MB 限制", 413)
    extension = PurePosixPath(filename.replace("\\", "/")).suffix.lower()
    warnings: list[str] = []
    if extension == ".csv":
        data = _csv_table(content, warnings)
    elif extension == ".xlsx":
        data = _xlsx_table(content, warnings)
    else:
        raise CleaningError("仅支持 CSV 和 XLSX；旧版 XLS 请另存为 XLSX 或 CSV")
    return Table(_filename(filename), data, warnings)


def inspect_table(content: bytes, filename: str) -> dict[str, Any]:
    table = read_table(content, filename)
    return {"filename": table.filename, "columns": list(table.data.columns), "total_rows": len(table.data),
            "preview": table.data.head(PREVIEW_ROWS).to_dict(orient="records"), "warnings": table.warnings}


def _finite_numeric(value: str) -> Decimal | None:
    try:
        # Plain numeric syntax only; underscores, NaN/Inf and formulas are not numbers.
        if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", value):
            return None
        number = Decimal(value)
        return number if number.is_finite() else None
    except (InvalidOperation, ValueError):
        return None


def _filter_mask(data: pd.DataFrame, rule: RowFilter) -> pd.Series:
    if rule.column not in data.columns:
        raise CleaningError("筛选条件引用了不存在的来源列")
    values = data[rule.column]
    if rule.operator in ("gt", "gte", "lt", "lte"):
        threshold = _finite_numeric(rule.value.strip())
        if threshold is None:
            raise CleaningError("数值筛选的比较值必须是有限数字")
        comparators = {"gt": lambda a: a > threshold, "gte": lambda a: a >= threshold,
                       "lt": lambda a: a < threshold, "lte": lambda a: a <= threshold}
        def compare(value: str) -> bool:
            number = _finite_numeric(value.strip())
            return number is not None and comparators[rule.operator](number)
        return values.map(compare)
    if rule.operator == "eq":
        return values.eq(rule.value)
    if rule.operator == "ne":
        return values.ne(rule.value)
    if rule.operator == "contains":
        return values.str.contains(rule.value, regex=False, na=False)
    empty = values.str.strip().eq("")
    return empty if rule.operator == "is_empty" else ~empty


def _csv_formula_needs_protection(value: str) -> bool:
    candidate = re.sub(r"^[\s\ufeff]+", "", value)
    dangerous = candidate.startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n"))
    return dangerous and not (candidate == value and _finite_numeric(value) is not None)


def _safe_csv_cells(data: pd.DataFrame) -> tuple[pd.DataFrame, int]:
    safe = data.copy()
    protected = 0
    def protect(value: str) -> str:
        nonlocal protected
        if _csv_formula_needs_protection(value):
            protected += 1
            return "'" + value
        return value
    for column in safe.columns:
        safe[column] = safe[column].map(protect)
    safe.columns = [protect(str(column)) for column in safe.columns]
    if len(safe.columns) != len(set(safe.columns)):
        raise CleaningError("公式防护后输出列名重名，请重命名冲突列")
    return safe, protected


class _BoundedCsvSink:
    def __init__(self):
        self.buffer = BytesIO()
        self.buffer.write(b"\xef\xbb\xbf")

    def write(self, value: str) -> int:
        chunk = value.encode("utf-8")
        if self.buffer.tell() + len(chunk) > MAX_OUTPUT_BYTES:
            raise CleaningError("下载 CSV 超过 80 MiB 输出预算，请减少行列或分批导出")
        return self.buffer.write(chunk)


def _csv_payload(data: pd.DataFrame) -> bytes:
    output = _BoundedCsvSink()
    writer = csv.writer(output, lineterminator="\r\n")
    writer.writerow(list(data.columns))
    writer.writerows(data.itertuples(index=False, name=None))
    return output.buffer.getvalue()


def _check_system_rows(data: pd.DataFrame, config: dict[str, Any], missing: list[str],
                       require_complete: bool) -> tuple[dict[str, Any], pd.Series, list[str]]:
    """Check format and current cohort eligibility without exposing raw records."""
    quality = {"checked": False, "checked_rows": 0, "valid_rows": 0, "invalid_rows": 0,
               "dropped_rows": 0, "remaining_rows": len(data), "complete_case_required": require_complete,
               "reasons": [], "invalid_preview": [], "column_errors": list(missing),
               "reason_counts_overlap": True, "category_warnings": [], "nonresponse_rows": 0}
    valid = pd.Series(True, index=data.index, dtype=bool)
    if missing:
        return quality, valid, ["必需列尚未齐全，未执行系统行检查；不能用剔除记录掩盖缺列。"]
    if config["id"] != "blood_pressure":
        return quality, valid, ["该结局尚无已实现的系统行质量规则，不能标记为系统就绪。"]
    bp_columns = [column for columns in config["outcome_columns"].values() for column in columns]
    descriptive_columns = [item["column"] for item in config["exposures"] + config["confounders"]]
    numeric_columns = bp_columns + descriptive_columns
    age_column = next(item["column"] for item in config["confounders"] if item["key"] == "age")
    numeric = _as_numeric(data, numeric_columns)
    finite = pd.DataFrame(np.isfinite(numeric), index=data.index, columns=numeric.columns)
    nonempty = pd.DataFrame({column: data[column].str.strip().ne("") for column in numeric_columns})
    invalid = pd.Series(False, index=data.index, dtype=bool)
    masks: list[tuple[str, pd.Series]] = []
    reasons = []
    def add(code: str, label: str, mask: pd.Series, columns: list[str]):
        nonlocal invalid
        count = int(mask.sum())
        if count:
            reasons.append({"code": code, "label": label, "count": count, "columns": columns})
            masks.append((code, mask))
            invalid = invalid | mask
    identifier = config["identifier_column"]
    add("identifier_missing", "参与者 ID 为空", data[identifier].str.strip().eq(""), [identifier])
    add("age_invalid", "年龄缺失或不是有限数字", ~finite[age_column], [age_column])
    add("age_under_18", "年龄不符合当前分析的成人队列（至少 18 岁）",
        finite[age_column] & numeric[age_column].lt(_SYSTEM_MIN_ADULT_AGE), [age_column])
    malformed_bp = nonempty[bp_columns] & (~finite[bp_columns] | numeric[bp_columns].le(0))
    add("blood_pressure_format_invalid", "非空血压读数含非数值、非有限数或非正数",
        malformed_bp.any(axis=1), [column for column in bp_columns if bool(malformed_bp[column].any())])
    ignored = pd.Series(False, index=data.index, dtype=bool)
    for kind, columns in config["outcome_columns"].items():
        lower, upper = _SYSTEM_BP_RANGES[kind]
        in_range = finite[columns] & numeric[columns].ge(lower) & numeric[columns].le(upper)
        add(f"{kind}_insufficient_readings", f"{'收缩压' if kind == 'sbp' else '舒张压'}在当前分析读数范围 {lower}–{upper} mmHg 内不足 2 次",
            in_range.sum(axis=1).lt(2), columns)
        ignored = ignored | (nonempty[columns] & finite[columns] & numeric[columns].gt(0) & ~in_range).any(axis=1)
    remaining_numeric = [column for column in descriptive_columns if column != age_column]
    malformed_descriptive = nonempty[remaining_numeric] & ~finite[remaining_numeric]
    add("numeric_field_invalid", "非空暴露/混杂字段不是有限数字（不自动猜测编码或单位）",
        malformed_descriptive.any(axis=1), [column for column in remaining_numeric if bool(malformed_descriptive[column].any())])
    warnings = []
    explicit_nonresponse = pd.Series(False, index=data.index, dtype=bool)
    nonresponse_columns = []
    category_warnings = []
    category_labels = {item["column"]: item["label"] for item in config["confounders"]}
    for column, response_codes in _NHANES_RESPONSE_CATEGORIES.items():
        if column not in numeric.columns:
            continue
        nonresponse_codes = _NHANES_EXPLICIT_NONRESPONSE.get(column, set())
        is_nonresponse = finite[column] & numeric[column].isin(nonresponse_codes)
        is_unknown = finite[column] & ~numeric[column].isin(response_codes | nonresponse_codes)
        nonresponse_count = int(is_nonresponse.sum())
        unknown_count = int(is_unknown.sum())
        if nonresponse_count or unknown_count:
            label = category_labels.get(column, column)
            category_warnings.append({"column": column, "label": label,
                                      "nonresponse_count": nonresponse_count, "unknown_count": unknown_count})
            if nonresponse_count:
                treatment = "已按字段完整选项标记为不合格" if require_complete else "默认保留原值；字段完整选项可将其标记为不合格"
                warnings.append(f"{label}有 {nonresponse_count} 条官方明确非回答编码（拒答/不知道），{treatment}，不自动重编码。")
            if unknown_count:
                warnings.append(f"{label}有 {unknown_count} 条未知有限分类编码，保留原值并提示核对；不自动映射、不因未知编码剔除。")
        explicit_nonresponse = explicit_nonresponse | is_nonresponse
        if nonresponse_count:
            nonresponse_columns.append(column)
    if require_complete:
        incomplete = ~nonempty[descriptive_columns]
        add("complete_case_missing", "已选择字段完整且排除明确非回答：暴露/混杂字段存在空白",
            incomplete.any(axis=1), [column for column in descriptive_columns if bool(incomplete[column].any())])
        add("explicit_nonresponse_code", "已选择字段完整且排除明确非回答：含官方拒答/不知道编码",
            explicit_nonresponse, nonresponse_columns)
    if bool(ignored.any()):
        warnings.append(f"{int(ignored.sum())} 条记录含正数但不在当前分析读数范围的血压值；这些读数不计入有效次数，其他每类至少 2 次有效读数时仍保留记录。")
    quality.update(checked=True, checked_rows=len(data), valid_rows=int((~invalid).sum()),
                   invalid_rows=int(invalid.sum()), reasons=reasons,
                   category_warnings=category_warnings, nonresponse_rows=int(explicit_nonresponse.sum()),
                   invalid_preview=[{"row_number": int(index) + 1,
                                     "reason_codes": [code for code, mask in masks if bool(mask.loc[index])]}
                                    for index in data.index[invalid][:PREVIEW_ROWS]])
    return quality, ~invalid, warnings


def _normalize_system_numeric_for_csv(data: pd.DataFrame, config: dict[str, Any]) -> tuple[pd.DataFrame, int, list[str]]:
    """Keep legitimate finite numeric values numeric after formula-safe export.

    With global trim disabled, e.g. ' +1' would otherwise be apostrophe-prefixed
    by formula protection. Strip only such dangerous-but-valid numeric cells,
    never ID/free-text cells. No rounding, coding, unit, or missing-value change.
    """
    columns = [column for group in config["outcome_columns"].values() for column in group]
    columns.extend(item["column"] for item in config["exposures"] + config["confounders"])
    columns = [column for column in columns if column in data.columns]
    numeric = _as_numeric(data, columns)
    changed_cells = 0
    changed_columns = []
    for column in columns:
        mask = np.isfinite(numeric[column]) & data[column].map(_csv_formula_needs_protection)
        count = int(mask.sum())
        if count:
            # Only whitespace is removed; leading '+' and numeric notation stay.
            data.loc[mask, column] = data.loc[mask, column].str.strip()
            changed_cells += count
            changed_columns.append(column)
    return data, changed_cells, changed_columns


def transform_table(content: bytes, filename: str, options: CleaningOptions) -> tuple[dict[str, Any], pd.DataFrame]:
    schema = get_cleaning_schema(options.outcome_id)
    table = read_table(content, filename)
    data = table.data.copy()
    audit: list[dict[str, Any]] = []
    warnings = list(table.warnings)
    def record(action: str, before: int, after: int, **extra):
        audit.append({"action": action, "before": before, "after": after, "removed": before - after, **extra})
    # Mapping/filter validity is checked even when earlier operations would empty the table.
    if any(item.source not in data.columns for item in options.columns):
        raise CleaningError("字段映射引用了不存在的来源列")
    if any(item.column not in data.columns for item in options.filters):
        raise CleaningError("筛选条件引用了不存在的来源列")
    if any(item.operator in ("gt", "gte", "lt", "lte") and _finite_numeric(item.value.strip()) is None
           for item in options.filters):
        raise CleaningError("数值筛选的比较值必须是有限数字")
    original_rows = len(data)
    if options.row_start > max(original_rows, 1) or (options.row_end is not None and options.row_end > original_rows):
        raise CleaningError("行范围超出原始数据记录数；范围按首行列名之后的 1 基数据记录编号")
    before = len(data)
    data = data.iloc[options.row_start - 1:options.row_end].copy()
    record("row_range", before, len(data), row_start=options.row_start, row_end=options.row_end)
    if options.trim_text:
        changed_cells = sum(data[column].ne(data[column].str.strip()).sum() for column in data.columns)
        for column in data.columns:
            data[column] = data[column].str.strip()
        record("trim_text", len(data), len(data), changed_cells=int(changed_cells))
    for index, rule in enumerate(options.filters, start=1):
        before = len(data)
        data = data.loc[_filter_mask(data, rule)].copy()
        record("filter", before, len(data), filter_index=index, column=rule.column, operator=rule.operator)
    if options.columns:
        before_columns = len(data.columns)
        data = data[[item.source for item in options.columns]].copy()
        data.columns = [item.target for item in options.columns]
        record("column_mapping", len(data), len(data), columns_before=before_columns, columns_after=len(data.columns))
    if options.drop_blank_rows:
        before = len(data)
        # Evaluate on selected OUTPUT columns, so irrelevant source columns do not keep a blank output row.
        blank = pd.DataFrame({column: data[column].str.strip().eq("") for column in data.columns}).all(axis=1)
        data = data.loc[~blank].copy()
        record("drop_blank_rows", before, len(data))
    if options.drop_duplicates:
        before = len(data)
        data = data.drop_duplicates(keep="first").copy()
        record("drop_duplicates", before, len(data))
        warnings.append("去重仅按当前输出列的完全相同值比较，保留首条；不按姓名或 ID 猜测同一人。")
    missing = [column for column in schema["required_columns"] if column not in data.columns]
    row_quality = None
    if options.mode == "system":
        row_quality, valid_rows, quality_warnings = _check_system_rows(
            data, OUTCOME_CONFIGS[options.outcome_id], missing, options.require_complete_system_rows)
        warnings.extend(quality_warnings)
        record("check_system_rows", len(data), len(data), checked=row_quality["checked"],
               invalid_rows=row_quality["invalid_rows"], complete_case_required=options.require_complete_system_rows,
               explicit_nonresponse_rows=row_quality["nonresponse_rows"])
        if options.drop_invalid_system_rows and row_quality["checked"]:
            before = len(data)
            data = data.loc[valid_rows].copy()
            row_quality["dropped_rows"] = before - len(data)
            row_quality["remaining_rows"] = len(data)
            record("drop_invalid_system_rows", before, len(data))
            if before != len(data):
                warnings.append(f"已按选择剔除 {before - len(data)} 条不合格记录；原文件保持不变，原因计数可能重叠。")
        elif row_quality["invalid_rows"]:
            warnings.append(f"有 {row_quality['invalid_rows']} 条不合格记录仍在输出中，请修正原始数据或选择剔除后重新预览。")
    data = data.reset_index(drop=True)
    ready = (options.mode == "system" and not missing and options.confirm_system_semantics and len(data) > 0
             and row_quality is not None and row_quality["checked"]
             and (row_quality["invalid_rows"] == 0 or options.drop_invalid_system_rows))
    if options.mode == "system":
        if missing:
            warnings.append(f"仍缺少 {len(missing)} 个系统必需列，请完成字段映射；不创建虚构的必需字段。")
        if not options.confirm_system_semantics:
            warnings.append("尚未确认血液样本、浓度单位和分类编码，不能作为系统格式导出。")
    else:
        warnings.append("自定义格式不会标记为系统就绪；用于群体导入前仍需满足系统列名和字段语义。")
    if len(data) == 0:
        warnings.append("清洗后没有数据记录，不能用于群体分析。")
    if options.mode == "system" and not missing:
        data, normalized_cells, normalized_columns = _normalize_system_numeric_for_csv(data, OUTCOME_CONFIGS[options.outcome_id])
        if normalized_cells:
            record("normalize_system_numeric_for_csv", len(data), len(data), changed_cells=normalized_cells,
                   columns=normalized_columns)
            warnings.append(f"为安全 CSV 导出，对 {normalized_cells} 个已确认为有限数字的系统数值单元格去除两端空白，避免公式防护将其变成文本；数值、单位和分类编码均未改变，ID/自由文本未规范化。")
    safe_data, protected = _safe_csv_cells(data)
    if protected:
        warnings.append(f"下载将为 {protected} 个可能触发电子表格公式的文本单元格/列名添加单引号前缀；合法负数保留。预览显示导出后的值。")
        # Show exactly the safe exported strings, not a preview that disagrees with download.
        if options.mode == "system" and list(safe_data.columns) != list(data.columns):
            ready = False
            warnings.append("列名触发公式防护，不能作为系统格式导出；请重命名该列。")
    preview = safe_data.head(PREVIEW_ROWS).to_dict(orient="records")
    preview_columns = list(safe_data.columns)
    record("csv_formula_protection", len(data), len(data), protected_cells=protected)
    result = {"filename": table.filename, "columns": preview_columns, "total_rows": len(data),
              "original_rows": original_rows, "preview": preview, "audit": audit, "warnings": warnings,
              "missing_required_columns": missing, "system_ready": bool(ready), "row_quality": row_quality}
    return result, safe_data


def preview_table(content: bytes, filename: str, options_json: str) -> dict[str, Any]:
    result, _ = transform_table(content, filename, parse_options(options_json))
    return result


def export_table(content: bytes, filename: str, options_json: str) -> tuple[bytes, str, dict[str, Any]]:
    options = parse_options(options_json)
    result, data = transform_table(content, filename, options)
    if options.mode == "system" and not result["system_ready"]:
        raise CleaningError("无法导出系统格式：请补齐必需列、确认样本/单位/编码、处理不合格记录，且清洗后至少保留一条合格记录")
    payload = _csv_payload(data)
    basename = PurePosixPath(_filename(filename)).stem
    return payload, f"cleaned_{basename}.csv", result
